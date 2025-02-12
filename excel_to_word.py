import pandas as pd
from docx import Document

def excel_to_word(excel_file, word_file):
    # Load the Excel file
    xls = pd.ExcelFile(excel_file)
    doc = Document()
    
    for sheet_name in xls.sheet_names:
        df = xls.parse(sheet_name)
        
        # Add sheet name as a heading
        doc.add_heading(sheet_name, level=1)
        
        # Convert dataframe to a Word table
        table = doc.add_table(rows=df.shape[0] + 1, cols=df.shape[1])
        table.style = 'Table Grid'
        
        # Add headers
        for j, col_name in enumerate(df.columns):
            table.cell(0, j).text = str(col_name)
        
        # Add data
        for i, row in df.iterrows():
            for j, value in enumerate(row):
                table.cell(i + 1, j).text = str(value)
        
        # Add a blank line between sheets
        doc.add_paragraph('\n')
    
    # Save the document
    doc.save(word_file)
    print(f'Excel file "{excel_file}" has been successfully converted to "{word_file}".')
